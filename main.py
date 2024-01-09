from tkinter import *
from tkinter.ttk import *
from tkinter.filedialog import askdirectory

import os
import openpyxl
from docx import Document

def get_title(doc_path):
    document = Document(doc_path)
    return [document.paragraphs[i].text for i in range(4)]


def write_excel_xlsx(path, value):
    index = len(value)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for i in range(0, index):
        for j in range(0, len(value[i])):
            sheet.cell(row=i+1, column=j+1, value=str(value[i][j]))
    workbook.save(path)


class WinGUI(Tk):
    def __init__(self):
        super().__init__()
        self.__win()
        self.tk_text_dir = self.__tk_text_lqop1jhv(self)
        self.tk_button_choose = self.__tk_button_lqop1uqh(self)
        self.tk_label_lqop3xk5 = self.__tk_label_lqop3xk5(self)
        self.tk_progressbar = self.__tk_progressbar_lqop49yg(self)
        self.tk_text_log = self.__tk_text_lqop4hx8(self)
        self.tk_button_process = self.__tk_button_lqop7bk2(self)
        
        self.tk_button_choose.configure(command=self.select_dir_path)
        self.tk_button_process.configure(command=self.process_dir)
        
        self.dir_path = None
        
        
    def select_dir_path(self):
        self.dir_path = askdirectory()
        self.tk_text_dir.delete(0.0, END)
        self.tk_text_dir.insert(END, self.dir_path)
    
    def set_progressbar(self, value):
        self.tk_progressbar['value'] = value
        self.update()
    
    def set_log(self, log):
        self.tk_text_log.insert(END, log+"\n")
        self.tk_text_log.see(END)
        self.update()
    
    def process_dir(self):
        if self.dir_path is None:
            self.set_log("请先选择文件夹")
            return
        if not os.path.exists(self.dir_path):
            self.set_log("文件夹不存在")
            return
        
        value = [['标题', '单位', '类型', '案号']]
        file_list = [file for file in os.listdir(self.dir_path) if (file.endswith('.docx')) and (not file.startswith('~$'))]
        for file in file_list:
            doc_path = os.path.join(self.dir_path, file)
            title = get_title(doc_path)
            value.append(title)
            self.set_log(f"{file} {title}")
            self.set_progressbar((len(value)-1)/len(file_list)*100)
        
        self.set_log("处理完成，请选择保存位置")
        save_dir = askdirectory()
        file_path = os.path.join(save_dir, 'export.xlsx')
        write_excel_xlsx(file_path, value)
        os.startfile(file_path)
    
    def __win(self):
        self.title("Word标题提取")
        # 设置窗口大小、居中
        width = 538
        height = 288
        screenwidth = self.winfo_screenwidth()
        screenheight = self.winfo_screenheight()
        geometry = '%dx%d+%d+%d' % (width, height, (screenwidth - width) / 2, (screenheight - height) / 2)
        self.geometry(geometry)
        
        self.resizable(width=False, height=False)
        
    def scrollbar_autohide(self,vbar, hbar, widget):
        """自动隐藏滚动条"""
        def show():
            if vbar: vbar.lift(widget)
            if hbar: hbar.lift(widget)
        def hide():
            if vbar: vbar.lower(widget)
            if hbar: hbar.lower(widget)
        hide()
        widget.bind("<Enter>", lambda e: show())
        if vbar: vbar.bind("<Enter>", lambda e: show())
        if vbar: vbar.bind("<Leave>", lambda e: hide())
        if hbar: hbar.bind("<Enter>", lambda e: show())
        if hbar: hbar.bind("<Leave>", lambda e: hide())
        widget.bind("<Leave>", lambda e: hide())
    
    def v_scrollbar(self,vbar, widget, x, y, w, h, pw, ph):
        widget.configure(yscrollcommand=vbar.set)
        vbar.config(command=widget.yview)
        vbar.place(relx=(w + x) / pw, rely=y / ph, relheight=h / ph, anchor='ne')
    def h_scrollbar(self,hbar, widget, x, y, w, h, pw, ph):
        widget.configure(xscrollcommand=hbar.set)
        hbar.config(command=widget.xview)
        hbar.place(relx=x / pw, rely=(y + h) / ph, relwidth=w / pw, anchor='sw')
    def create_bar(self,master, widget,is_vbar,is_hbar, x, y, w, h, pw, ph):
        vbar, hbar = None, None
        if is_vbar:
            vbar = Scrollbar(master)
            self.v_scrollbar(vbar, widget, x, y, w, h, pw, ph)
        if is_hbar:
            hbar = Scrollbar(master, orient="horizontal")
            self.h_scrollbar(hbar, widget, x, y, w, h, pw, ph)
        self.scrollbar_autohide(vbar, hbar, widget)
    
    def __tk_text_lqop1jhv(self,parent):
        text = Text(parent)
        text.place(x=20, y=60, width=350, height=30)
        return text
    
    def __tk_button_lqop1uqh(self,parent):
        btn = Button(parent, text="选择文件夹", takefocus=False)
        btn.place(x=400, y=60, width=100, height=30)
        return btn
    
    def __tk_label_lqop3xk5(self,parent):
        label = Label(parent,text="标签",anchor="center", )
        label.place(x=24, y=4, width=480, height=42)
        return label
    
    def __tk_progressbar_lqop49yg(self,parent):
        progressbar = Progressbar(parent, orient=HORIZONTAL,)
        progressbar.place(x=20, y=110, width=350, height=30)
        return progressbar
    
    def __tk_text_lqop4hx8(self,parent):
        text = Text(parent)
        text.place(x=20, y=160, width=350, height=100)
        return text
    
    def __tk_button_lqop7bk2(self,parent):
        btn = Button(parent, text="开始处理", takefocus=False,)
        btn.place(x=400, y=110, width=100, height=30)
        return btn
    
class Win(WinGUI):
    def __init__(self, controller):
        self.ctl = controller
        super().__init__()
        self.__event_bind()
        self.ctl.init(self)
    def __event_bind(self):
        pass
    
if __name__ == "__main__":
    win = WinGUI()
    win.mainloop()